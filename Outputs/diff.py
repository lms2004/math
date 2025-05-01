import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

# 设置路径
root = "./Outputs/"
better_path = "test_better_predictions.xlsx"
original_path = "test_original_predictions.xlsx"

# 读取数据
better = pd.read_excel(root + better_path)
original = pd.read_excel(root + original_path)

# 校验数据一致性
if not better.columns.equals(original.columns):
    raise ValueError("列名不一致")
if len(better) != len(original):
    raise ValueError("数据行数不一致")

# 标记差异行（仅比较predicted_label）
diff_mask = better['predicted_label'] != original['predicted_label']
diff_indices = diff_mask[diff_mask].index

if len(diff_indices) == 0:
    print("预测结果完全一致")
else:
    # 合并差异数据
    merged = pd.concat([
        better.loc[diff_indices].add_suffix('_Better'), 
        original.loc[diff_indices].add_suffix('_Original')
    ], axis=1)
    
    # 调整列顺序（特征列并排显示）
    columns_order = []
    for col in better.columns:
        columns_order.append(f"{col}_Better")
        columns_order.append(f"{col}_Original")
    merged = merged[columns_order]
    
    # 输出到Excel
    output_path = root + "prediction_differences.xlsx"
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        merged.to_excel(writer, index=False)
        
        # 获取工作表对象
        workbook = writer.book
        worksheet = writer.sheets['Sheet1']
        
        # 设置颜色格式
        green_fill = PatternFill(start_color='90EE90', end_color='90EE90', fill_type='solid')
        red_fill = PatternFill(start_color='FFB6C1', end_color='FFB6C1', fill_type='solid')
        
        # 自动调整列宽（限制最大宽度为50字符）
        for column in worksheet.columns:
            max_length = 0
            column = [cell for cell in column]
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[get_column_letter(column[0].column)].width = adjusted_width
        
        # 标注预测列颜色
        pred_better_col = merged.columns.get_loc('predicted_label_Better') + 1  # Excel从1开始
        pred_original_col = merged.columns.get_loc('predicted_label_Original') + 1
        
        for row in range(2, len(merged)+2):  # 从数据行开始
            worksheet.cell(row=row, column=pred_better_col).fill = green_fill
            worksheet.cell(row=row, column=pred_original_col).fill = red_fill
        
        # 冻结首行
        worksheet.freeze_panes = 'A2'
        
        # 添加筛选器
        worksheet.auto_filter.ref = worksheet.dimensions

    print(f"差异文件已生成：{output_path}")
    print(f"共发现 {len(diff_indices)} 处差异（绿色为改进模型，红色为原始模型）")